#!/usr/bin/env python3
# Copyright (c) Qualcomm Technologies, Inc. and/or its subsidiaries.
# SPDX-License-Identifier: BSD-3-Clause
# =============================================================================
# Skill        : nhlos-carveout-validation
# Display Name : NHLOS Carve-out Reporter
# Version      : 0.0.1
# Platform     : Qualcomm Linux
# Description  : Parses NHLOS carve-outs from device tree reserved-memory
#                nodes and reports their layout and sizes.
#                No external spec or IP-XACT file required.
#
# Usage:
#   python scripts/validate.py --dump <dump_dir>
#   python scripts/validate.py -d mem_dump/ --report-dir reports/
# =============================================================================

import argparse
import datetime
import json
import sys
from pathlib import Path


# ---------------------------------------------------------------------------
# Device tree reserved-memory parsers (two formats supported)
# ---------------------------------------------------------------------------

def parse_dt_reserved_memory(res_dir: Path) -> list[dict]:
    """Parse device tree reserved-memory binary nodes (memory-map skill format).

    Each .bin file represents one reserved-memory node.
    Binary layout: repeated 16-byte records of (addr: 8 bytes BE, size: 8 bytes BE).

    Returns list of dicts: name, base, size_bytes, size_kb, size_mb.
    """
    regions: list[dict] = []
    for bin_file in sorted(res_dir.glob("*.bin")):
        data = bin_file.read_bytes()
        if len(data) < 16:
            continue
        total_bytes = 0
        base_addr   = None
        for i in range(0, len(data) - 15, 16):
            addr = int.from_bytes(data[i:i+8],   byteorder="big")
            size = int.from_bytes(data[i+8:i+16], byteorder="big")
            if base_addr is None:
                base_addr = addr
            total_bytes += size
        if total_bytes > 0 and base_addr is not None:
            regions.append({
                "name":       bin_file.stem,
                "base":       base_addr,
                "size_bytes": total_bytes,
                "size_kb":    total_bytes // 1024,
                "size_mb":    round(total_bytes / 1024 / 1024, 4),
            })
    return sorted(regions, key=lambda r: r["base"])


def parse_dma_reservations(txt_path: Path) -> list[dict]:
    """Parse DMA_reservations.txt (reference script format).

    Format per line:
      name  high_base  low_base  high_size  low_size
    All address/size fields are 8-digit hex without 0x prefix.

    Returns list of dicts: name, base, size_bytes, size_kb, size_mb.
    """
    regions: list[dict] = []
    try:
        for line in txt_path.read_text(errors="replace").splitlines():
            line = line.strip()
            if not line or "No such file" in line or "xxd: command not found" in line:
                continue
            if "disabled" in line.lower():
                continue
            parts = line.split()
            if len(parts) < 5:
                continue
            try:
                name       = parts[0]
                high_base  = int(parts[1], 16)
                low_base   = int(parts[2], 16)
                high_size  = int(parts[3], 16)
                low_size   = int(parts[4], 16)
                base       = (high_base << 32) | low_base
                size_bytes = (high_size << 32) | low_size
                if size_bytes == 0:
                    continue
                regions.append({
                    "name":       name,
                    "base":       base,
                    "size_bytes": size_bytes,
                    "size_kb":    size_bytes // 1024,
                    "size_mb":    round(size_bytes / 1024 / 1024, 4),
                })
            except (ValueError, IndexError):
                pass
    except Exception:
        pass
    return sorted(regions, key=lambda r: r["base"])


def load_carveouts(dump_dir: Path) -> tuple[list[dict], str]:
    """Auto-detect and load carve-outs from a dump directory.

    Supports two formats:
      1. reserved-memory/*.bin  -- memory-map skill format
      2. DMA_reservations.txt   -- reference script format

    Returns (regions, format_name).
    """
    # Format 1: binary reserved-memory nodes
    res_dir = dump_dir / "reserved-memory"
    if res_dir.exists() and any(res_dir.glob("*.bin")):
        return parse_dt_reserved_memory(res_dir), "reserved-memory/*.bin"

    # Format 2: DMA_reservations.txt from reference scripts
    dma_txt = dump_dir / "DMA_reservations.txt"
    if dma_txt.exists():
        return parse_dma_reservations(dma_txt), "DMA_reservations.txt"

    return [], "none"


# ---------------------------------------------------------------------------
# Summary
# ---------------------------------------------------------------------------

def _summary(results: list[dict]) -> dict:
    total_kb = sum(r["actual_size_kb"] for r in results)
    return {
        "found":           len(results),
        "total_actual_mb": round(total_kb / 1024, 2),
    }


# ---------------------------------------------------------------------------
# Report rendering
# ---------------------------------------------------------------------------

_STATUS_COLOR_HTML = {
    "FOUND": "#d4edda",
}


def save_json(results: list[dict], dump_dir: Path, carveout_source: str,
              report_dir: Path) -> None:
    """Save structured CarveoutValidationReport JSON."""
    summ = _summary(results)
    report = {
        "report_type":      "CarveoutValidationReport",
        "generated_at":     datetime.datetime.now().isoformat(),
        "dump_dir":         str(dump_dir),
        "carveout_source":  carveout_source,
        "results":          results,
        "summary":          summ,
    }
    path = report_dir / "carveout_validation_report.json"
    path.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(f"  JSON saved : {path}")


def save_txt(results: list[dict], dump_dir: Path, carveout_source: str,
             report_dir: Path) -> None:
    SEP  = "=" * 80
    SEP2 = "-" * 80
    summ = _summary(results)
    lines = [
        SEP,
        "  NHLOS CARVE-OUT REPORT",
        SEP,
        f"  Source   : {carveout_source}",
        f"  Dump     : {dump_dir}",
        f"  Captured : {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        SEP,
        f"  {'Name':<30} {'Base Address':<16} {'Size (MB)':>12}  Status",
        SEP2,
    ]

    for r in results:
        base_str = hex(r["base_addr_int"]) if r.get("base_addr_int") else r.get("base_addr", "--")
        lines.append(
            f"  {r['name']:<30} {base_str:<16} {r['actual_size_mb']:>10.2f} MB  [OK] FOUND"
        )

    lines += [
        SEP,
        "  SUMMARY",
        SEP2,
        f"  Carve-outs found : {summ['found']}",
        f"  Total size       : {summ['total_actual_mb']:.2f} MB",
        SEP,
        "",
    ]

    path = report_dir / "nhlos_validation_report.txt"
    path.write_text("\n".join(lines), encoding="utf-8")
    print(f"  TXT  saved : {path}")


def save_html(results: list[dict], dump_dir: Path, carveout_source: str,
              report_dir: Path) -> None:
    summ = _summary(results)

    rows_html = ""
    for r in results:
        base_s = r.get("base_addr", "--")
        rows_html += (
            f'<tr style="background:#d4edda">'
            f'<td>{r["name"]}</td>'
            f'<td class="mono">{base_s}</td>'
            f'<td class="num">{r["actual_size_mb"]:.2f} MB</td>'
            f'<td>[OK] FOUND</td>'
            f'</tr>\n'
        )

    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>NHLOS Carve-out Report</title>
<style>
  body  {{ font-family: Segoe UI, Arial, sans-serif; font-size: 13px;
           background: #f5f5f5; margin: 0; padding: 24px; color: #222; }}
  h1    {{ font-size: 20px; margin-bottom: 4px; }}
  h2    {{ font-size: 15px; margin: 20px 0 6px; border-bottom: 2px solid #ccc;
           padding-bottom: 4px; }}
  .meta {{ font-size: 12px; color: #555; margin-bottom: 16px; }}
  table {{ border-collapse: collapse; width: 100%; background: #fff;
           border-radius: 6px; box-shadow: 0 1px 4px rgba(0,0,0,.12);
           margin-bottom: 24px; }}
  th    {{ background: #3a3a3a; color: #fff; padding: 7px 10px;
           text-align: left; font-weight: 600; font-size: 12px; }}
  td    {{ padding: 5px 10px; border-bottom: 1px solid #eee; font-size: 12px; }}
  tr:last-child td {{ border-bottom: none; }}
  .num  {{ text-align: right; font-variant-numeric: tabular-nums; white-space: nowrap; }}
  .mono {{ font-family: monospace; }}
</style>
</head>
<body>
<h1>NHLOS Carve-out Report</h1>
<div class="meta">
  <strong>Source:</strong> {carveout_source} &nbsp;&nbsp;
  <strong>Dump:</strong> {dump_dir} &nbsp;&nbsp;
  <strong>Captured:</strong> {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}
</div>

<h2>Carve-out Layout ({summ['found']} regions, {summ['total_actual_mb']:.2f} MB total)</h2>
<table>
  <tr>
    <th>Name</th><th>Base Address</th>
    <th class="num">Size (MB)</th><th>Status</th>
  </tr>
  {rows_html}
</table>

<h2>Summary</h2>
<table style="max-width:320px">
  <tr><th>Metric</th><th class="num">Value</th></tr>
  <tr style="background:#d4edda">
    <td>Carve-outs found</td>
    <td class="num">{summ['found']}</td>
  </tr>
  <tr style="background:#1a5276;color:#fff;font-weight:700">
    <td>Total Size</td>
    <td class="num">{summ['total_actual_mb']:.2f} MB</td>
  </tr>
</table>
</body>
</html>
"""
    path = report_dir / "nhlos_validation_report.html"
    path.write_text(html, encoding="utf-8")
    print(f"  HTML saved : {path}")


def save_xlsx(results: list[dict], dump_dir: Path, carveout_source: str,
              report_dir: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("  [SKIP] openpyxl not installed; skipping Excel report.")
        return

    summ  = _summary(results)
    wb    = Workbook()
    _thin = Side(style="thin", color="CCCCCC")
    BDR   = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

    FOUND_FILL = PatternFill("solid", fgColor="D4EDDA")
    HDR_FILL   = PatternFill("solid", fgColor="3A3A3A")
    HDR_FONT   = Font(bold=True, color="FFFFFF", size=11)
    GRAND      = PatternFill("solid", fgColor="1A5276")

    # ---- Sheet 1: Carve-out Layout ----
    ws1 = wb.active
    ws1.title = "Carveout Layout"
    widths = [30, 16, 14, 14]
    for i, w in enumerate(widths, 1):
        ws1.column_dimensions[get_column_letter(i)].width = w

    ws1.append([f"Source: {carveout_source}   Dump: {dump_dir}   "
                f"Captured: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"])
    ws1.cell(1, 1).font = Font(bold=True, size=12)
    ws1.append([])

    headers = ["Name", "Base Address", "Size (MB)", "Status"]
    ws1.append(headers)
    r = ws1.max_row
    for col in range(1, len(headers) + 1):
        c = ws1.cell(r, col)
        c.font = HDR_FONT
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center")
        c.border = BDR

    for res in results:
        row_vals = [
            res["name"],
            res.get("base_addr", "--"),
            round(res["actual_size_mb"], 2),
            "FOUND",
        ]
        ws1.append(row_vals)
        r = ws1.max_row
        for col in range(1, len(headers) + 1):
            c = ws1.cell(r, col)
            c.border = BDR
            c.fill = FOUND_FILL
            if col == 3 and isinstance(c.value, float):
                c.alignment = Alignment(horizontal="right")
                c.number_format = "#,##0.00"

    # ---- Sheet 2: Summary ----
    ws2 = wb.create_sheet("Summary")
    ws2.column_dimensions["A"].width = 24
    ws2.column_dimensions["B"].width = 14

    ws2.append(["Metric", "Value"])
    r = ws2.max_row
    for col in range(1, 3):
        c = ws2.cell(r, col)
        c.font = HDR_FONT
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center")
        c.border = BDR

    ws2.append(["Carve-outs found", summ["found"]])
    r = ws2.max_row
    for col in range(1, 3):
        c = ws2.cell(r, col)
        c.border = BDR
        c.fill = FOUND_FILL

    ws2.append(["Total Size", f"{summ['total_actual_mb']:.2f} MB"])
    r = ws2.max_row
    for col in range(1, 3):
        c = ws2.cell(r, col)
        c.fill = GRAND
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.border = BDR

    path = report_dir / "nhlos_validation_report.xlsx"
    try:
        wb.save(str(path))
        print(f"  XLSX saved : {path}")
    except PermissionError:
        alt = report_dir / "nhlos_validation_report_new.xlsx"
        wb.save(str(alt))
        print(f"  XLSX saved : {alt}  (original is open)")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "NHLOS Carve-out Reporter v0.0.1\n\n"
            "Parses NHLOS carve-outs from device tree reserved-memory nodes\n"
            "and reports their layout and sizes.\n\n"
            "Example:\n"
            "  python scripts/validate.py -d mem_dump/\n"
            "  python scripts/validate.py -d mem_dump/ --report-dir reports/"
        ),
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument("--dump",       "-d", metavar="DIR",  required=True,
                        help="Dump directory containing reserved-memory/ binaries "
                             "or DMA_reservations.txt")
    parser.add_argument("--report-dir",       metavar="DIR",  default=".",
                        help="Output directory for reports (default: .)")
    args = parser.parse_args()

    dump_dir   = Path(args.dump)
    report_dir = Path(args.report_dir)
    report_dir.mkdir(parents=True, exist_ok=True)

    if not dump_dir.exists():
        sys.exit(f"[ERROR] Dump directory not found: {dump_dir}")

    print(f"[1/2] Reading carve-outs from {dump_dir}")
    raw_regions, fmt = load_carveouts(dump_dir)
    if not raw_regions:
        sys.exit(
            f"[ERROR] No carve-out data found in {dump_dir}.\n"
            f"  Expected one of:\n"
            f"    - reserved-memory/*.bin  (memory-map skill format)\n"
            f"    - DMA_reservations.txt   (reference script format)"
        )
    print(f"  Format        : {fmt}")
    print(f"  DT carve-outs : {len(raw_regions)}")

    # Build result records
    results = []
    for r in raw_regions:
        results.append({
            "name":            r["name"],
            "base_addr":       hex(r["base"]),
            "base_addr_int":   r["base"],
            "actual_size_kb":  r["size_kb"],
            "actual_size_mb":  r["size_mb"],
            "status":          "FOUND",
        })

    summ = _summary(results)
    print(f"\n  Carve-outs found : {summ['found']}")
    print(f"  Total size       : {summ['total_actual_mb']:.2f} MB")

    print(f"\n[2/2] Saving reports to {report_dir} ...")
    save_json(results, dump_dir, fmt, report_dir)
    save_txt(results,  dump_dir, fmt, report_dir)
    save_html(results, dump_dir, fmt, report_dir)
    save_xlsx(results, dump_dir, fmt, report_dir)
    print("\nDone.")


if __name__ == "__main__":
    main()